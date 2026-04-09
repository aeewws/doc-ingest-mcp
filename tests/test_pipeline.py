from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from doc_ingest_mcp import pipeline
from doc_ingest_mcp.backend import _extract_docx_text
from doc_ingest_mcp.mcp_server import build_tool_registry
from doc_ingest_mcp.pipeline import batch_ingest, export_directory, ingest_file
from doc_ingest_mcp.types import ChunkRecord, DocumentArtifact, IngestPaths, IngestResult, Manifest


def test_ingest_writes_expected_contract(tmp_path: Path, fixture_dir: Path) -> None:
    source = fixture_dir / "sample.pdf"
    result = ingest_file(source, output_dir=tmp_path / "sample", backend="fake")

    manifest_path = Path(result.paths.manifest_path)
    markdown_path = Path(result.paths.markdown_path)
    chunks_path = Path(result.paths.chunks_path)
    tables_dir = Path(result.paths.tables_dir)

    assert manifest_path.exists()
    assert markdown_path.exists()
    assert chunks_path.exists()
    assert tables_dir.exists()

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["source"].endswith("sample.pdf")
    assert manifest["status"] == "ok"
    assert manifest["chunk_count"] >= 1
    assert manifest["table_count"] == 1
    assert "Alpha paragraph." in markdown_path.read_text(encoding="utf-8")
    assert "Beta paragraph." in (tables_dir / "table-001.md").read_text(encoding="utf-8")


def test_batch_ingest_processes_three_sample_documents(tmp_path: Path, fixture_dir: Path) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    for name in ["sample.pdf", "scan.png", "contract.docx"]:
        (inbox / name).write_bytes((fixture_dir / name).read_bytes())

    results = batch_ingest(inbox, output_root=tmp_path / "out", backend="fake")
    assert len(results) == 3
    assert all(Path(result.paths.output_dir).exists() for result in results)


def test_export_reads_generated_output(tmp_path: Path, fixture_dir: Path) -> None:
    result = ingest_file(fixture_dir / "contract.docx", output_dir=tmp_path / "sample", backend="fake")
    payload = export_directory(Path(result.paths.output_dir), "json")
    data = json.loads(payload)
    assert data["manifest"]["status"] == "ok"
    assert data["chunks"]
    assert "Clause one." in Path(result.paths.markdown_path).read_text(encoding="utf-8")


def test_tool_registry_names_are_stable() -> None:
    registry = build_tool_registry()
    assert list(registry) == [
        "ingest_document",
        "batch_ingest",
        "get_manifest",
        "read_chunks",
        "export_markdown",
    ]


def test_extract_docx_text_keeps_paragraphs_and_tables(tmp_path: Path) -> None:
    source = tmp_path / "mixed.docx"
    document = Document()
    document.add_paragraph("Intro paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Column A"
    table.rows[0].cells[1].text = "Column B"
    table.rows[1].cells[0].text = "Value 1"
    table.rows[1].cells[1].text = "Value 2"
    document.save(source)

    text = _extract_docx_text(source)

    assert "Intro paragraph." in text
    assert "| Column A | Column B |" in text
    assert "| --- | --- |" in text
    assert "| Value 1 | Value 2 |" in text


def test_watch_inbox_skips_failures_and_continues_processing(tmp_path: Path, monkeypatch) -> None:
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    bad_source = inbox / "bad.pdf"
    good_source = inbox / "good.pdf"
    bad_source.write_text("bad", encoding="utf-8")
    good_source.write_text("good", encoding="utf-8")

    def fake_ingest_file(source: Path, output_dir: Path | None = None, backend: str | None = None) -> IngestResult:
        if source.name == "bad.pdf":
            raise RuntimeError("boom")
        return IngestResult(
            artifact=DocumentArtifact(
                manifest=Manifest(
                    source=str(source),
                    mime_type="application/pdf",
                    status="ok"
                ),
                markdown="ok",
                chunks=[ChunkRecord(id="chunk-1", text="ok")]
            ),
            paths=IngestPaths(
                output_dir=str(output_dir or source.parent / "out" / source.stem),
                manifest_path="manifest.json",
                markdown_path="document.md",
                chunks_path="chunks.json",
                assets_dir="assets",
                tables_dir="tables"
            )
        )

    monkeypatch.setattr(pipeline, "ingest_file", fake_ingest_file)

    results = pipeline.watch_inbox(inbox, output_root=tmp_path / "out", backend="fake", once=True)

    assert len(results) == 1
    assert results[0].artifact.manifest.source.endswith("good.pdf")
