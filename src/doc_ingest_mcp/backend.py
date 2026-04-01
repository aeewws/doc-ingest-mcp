from __future__ import annotations

from pathlib import Path

from .errors import MissingDependencyError, UnsupportedInputError
from .types import BackendDocument

SUPPORTED_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg", ".docx"}


def _read_text_fallback(source: Path) -> str:
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf_text(source)
        if text:
            return text
    elif suffix == ".docx":
        text = _extract_docx_text(source)
        if text:
            return text

    raw = source.read_bytes()
    text = raw.decode("utf-8", errors="ignore").strip()
    return text or source.stem


def _extract_pdf_text(source: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception:
        return ""

    try:
        reader = PdfReader(str(source))
    except Exception:
        return ""

    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(page.strip() for page in pages if page.strip()).strip()
    return text


def _extract_docx_text(source: Path) -> str:
    try:
        from docx import Document
    except Exception:
        return ""

    try:
        document = Document(str(source))
    except Exception:
        return ""

    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    text = "\n\n".join(paragraphs).strip()
    if text:
        return text

    table_blocks: list[str] = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("|", r"\|") for cell in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
        if rows:
            header = rows[0]
            separator = "| " + " | ".join(["---"] * header.count("|")) + " |"
            table_blocks.append("\n".join([header, separator, *rows[1:]]))
    return "\n\n".join(table_blocks).strip()


def _fake_tables(text: str) -> list[str]:
    blocks = [block.strip() for block in text.split("\n\n") if block.strip()]
    if len(blocks) <= 1:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if len(lines) > 1:
            blocks = lines[:3]
    if len(blocks) > 1:
        header = "| section | text |\n| --- | --- |"
        rows = []
        for index, block in enumerate(blocks[:3]):
            escaped = block.replace("|", r"\|")
            rows.append(f"| {index + 1} | {escaped} |")
        return [header + "\n" + "\n".join(rows)]
    return []


def _require_supported_source(source: Path) -> None:
    if source.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise UnsupportedInputError(
            f"Unsupported input type: {source.suffix or '<none>'}. "
            "Supported inputs are pdf, png, jpg, jpeg, and docx."
        )


def _extract_docling_document(source: Path) -> BackendDocument:
    try:
        from docling.document_converter import DocumentConverter
    except Exception as exc:  # pragma: no cover - dependency path
        raise MissingDependencyError(
            "docling",
            "Install the optional parser with `pip install -e \".[docling]\"` "
            "or use `--backend fake` for smoke tests.",
        ) from exc

    converter = DocumentConverter()
    result = converter.convert(str(source))
    document = getattr(result, "document", result)

    markdown = _call_first_text(document, ["export_to_markdown", "to_markdown", "export_markdown"])
    if not markdown:
        markdown = _call_first_text(result, ["export_to_markdown", "to_markdown"])
    if not markdown:
        markdown = _read_text_fallback(source)

    page_count = _infer_page_count(result, document)
    tables = _collect_tables(result, document)
    asset_paths = _collect_assets(result, document)
    ocr_used = bool(getattr(result, "ocr_used", None) or getattr(document, "ocr_used", None))
    return BackendDocument(
        source=source,
        markdown=markdown,
        page_count=page_count,
        ocr_used=ocr_used,
        table_markdowns=tables,
        asset_paths=asset_paths,
    )


def _collect_assets(result: object, document: object) -> list[Path]:
    for candidate in (getattr(result, "assets", None), getattr(document, "assets", None)):
        if not candidate:
            continue
        items: list[Path] = []
        for item in candidate:
            path = getattr(item, "path", item)
            try:
                items.append(Path(path))
            except TypeError:
                continue
        if items:
            return items
    return []


def _collect_tables(result: object, document: object) -> list[str]:
    for candidate in (getattr(result, "tables", None), getattr(document, "tables", None)):
        if not candidate:
            continue
        tables: list[str] = []
        for table in candidate:
            text = _call_first_text(table, ["export_to_markdown", "to_markdown", "export_to_text"])
            if text:
                tables.append(text)
        if tables:
            return tables
    return []


def _infer_page_count(result: object, document: object) -> int:
    for candidate in (
        getattr(result, "pages", None),
        getattr(document, "pages", None),
        getattr(result, "page_count", None),
        getattr(document, "page_count", None),
    ):
        if candidate is None:
            continue
        if isinstance(candidate, int) and candidate > 0:
            return candidate
        try:
            return max(1, len(candidate))
        except TypeError:
            continue
    return 1


def _call_first_text(candidate: object, method_names: list[str]) -> str:
    for name in method_names:
        method = getattr(candidate, name, None)
        if not callable(method):
            continue
        try:
            value = method()
        except TypeError:
            try:
                value = method("markdown")
            except Exception:
                continue
        except Exception:
            continue
        if value is None:
            continue
        text = value if isinstance(value, str) else str(value)
        if text.strip():
            return text
    if isinstance(candidate, str) and candidate.strip():
        return candidate
    return ""


def extract_document(source: Path, backend: str = "auto") -> BackendDocument:
    _require_supported_source(source)
    mode = (backend or "auto").strip().lower()
    if mode not in {"auto", "fake", "docling"}:
        raise UnsupportedInputError(f"Unknown backend '{backend}'. Choose auto, fake, or docling.")
    if mode == "fake":
        text = _read_text_fallback(source)
        return BackendDocument(
            source=source,
            markdown=text,
            page_count=max(1, text.count("\f") + 1),
            ocr_used=False,
            table_markdowns=_fake_tables(text),
            asset_paths=[],
        )
    return _extract_docling_document(source)
